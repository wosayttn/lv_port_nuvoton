import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import os

# 讀取 XiP.xlsx
xlsx_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'XiP.xlsx')
wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]

# 解析數據：只取 "All scenes avg." 列
groups = []  # 儲存各組名稱
avg_cpu = []
avg_fps = []
for row in ws.iter_rows(min_row=2, values_only=True):
    place, width_val, height_val, bpp, name, cpu, fps, *_ = row
    if place is None or name is None:
        break
    if name == 'All scenes avg.':
        groups.append(place)
        avg_cpu.append(int(round(cpu * 100)))
        avg_fps.append(fps)

wb.close()

x = np.arange(len(groups))  # 項目位置
width = 0.5  # 每個直條的寬度

# 顏色配置
colors = ['#d4af37', '#1f77b4', '#2ca02c', '#d62728', '#9467bd', '#8c564b']

# 建立圖表
fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(8, 10), sharex=True)

# --- 1. Avg. FPS 直條圖 ---
ax1.bar(x, avg_fps, width, color=[colors[i % len(colors)] for i in range(len(groups))])
ax1.set_ylabel('Avg. FPS', fontsize=12)
ax1.set_title('All Scenes Avg. FPS', fontsize=14, fontweight='bold')
ax1.grid(True, linestyle='--', alpha=0.5, axis='y')

# --- 2. Avg. CPU 直條圖 ---
ax2.bar(x, avg_cpu, width, color=[colors[i % len(colors)] for i in range(len(groups))])
ax2.set_ylabel('Avg. CPU (%)', fontsize=12)
ax2.set_title('All Scenes Avg. CPU', fontsize=14, fontweight='bold')
ax2.grid(True, linestyle='--', alpha=0.5, axis='y')

# 設定 X 軸標籤
ax2.set_xticks(x)
ax2.set_xticklabels(groups, rotation=0, ha='center')

plt.tight_layout()

# 顯示圖表
plt.show()